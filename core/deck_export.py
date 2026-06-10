"""10-slide executive PPTX built from live run summary + ROI scenario."""
import io
from pptx import Presentation
from pptx.util import Pt

TITLE = "DataEng Copilot — Claude for Data Engineering"


def _slide(prs, title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    body = s.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = b
        p.font.size = Pt(18)
    return s


def build_deck(summary: dict, roi_data: dict) -> bytes:
    prs = Presentation()
    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = TITLE
    t.placeholders[1].text = ("Same productivity. Fraction of the headcount cost.\n"
                              "Built on Claude (claude-opus-4-8).")

    _slide(prs, "The Problem", [
        "Data engineering teams drown in repetitive lifecycle work",
        "Profiling, ETL, quality triage, ad-hoc SQL, reporting — every week",
        "Backlogs grow faster than headcount budgets"])
    _slide(prs, "The Solution", [
        "An autonomous Claude agent runs the full lifecycle",
        "Ingest → profile → clean → anomaly scan → answer → report",
        "Human engineers review and steer instead of typing SQL"])
    _slide(prs, "Architecture", [
        "Claude (claude-opus-4-8) manual tool-use loop",
        "Local tools: DuckDB SQL, profiling, 5-layer statistical anomaly detector",
        "Every step logged: tokens, cost, latency — fully auditable",
        "Replay mode = deterministic demos, zero API risk"])
    _slide(prs, "Live Demo Results", [
        f"{summary['tool_calls']} autonomous tool calls",
        f"Agent time: {summary['agent_minutes']} min — human baseline: {summary['human_hours']} h",
        f"Speedup: {summary['speedup_x']}x",
        f"API cost: ${summary['total_cost_usd']}",
        f"Tokens: {summary['tokens_in']:,} in / {summary['tokens_out']:,} out"])
    _slide(prs, "What It Found", [
        "Seeded data-quality issues: fixed autonomously",
        "Revenue anomalies: flagged with week, severity, magnitude",
        "Business questions: answered with charts, exec-ready report"])
    _slide(prs, "ROI Scenario (illustrative)", [
        f"FTE-equivalent capacity freed: {roi_data['fte_freed']}",
        f"Annual savings at full ramp: ${roi_data['annual_savings']:,.0f}",
        f"Payback month: {roi_data['payback_month']}",
        "Framing choice: reduce headcount OR multiply output — same math"])
    _slide(prs, "Security & Operations", [
        "Runs on your data, your infra — agent tools execute locally",
        "Full event audit trail per run",
        "Human-in-the-loop review before anything ships"])
    _slide(prs, "Roadmap", [
        "Pilot: one team, 4 weeks, measure baseline vs agent",
        "Expand: orchestration (Airflow/dbt), warehouse connectors",
        "Scale: agent fleet across all data domains"])
    _slide(prs, "The Ask", [
        "30-minute live demo with your data team",
        "Partnership conversation with Anthropic",
        "Contact: gsundar4121@gmail.com"])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
